"""Generates the per-template reference .docx files used by
rendering/renderer.py (pypandoc's --reference-doc). Pandoc reuses whatever
Heading 1 / Heading 2 / Normal paragraph and font styles it finds in the
named reference document, so editing a template's look only requires
re-running this script — no changes to the runtime rendering code.

Run manually whenever a template's style should change:
    python -m rendering.assets.docx_templates.generate_reference_docs
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

_OUTPUT_DIR = Path(__file__).parent


def _set_style(document: Document, style_name: str, font_name: str, size: int, color: str, bold: bool):
    style = document.styles[style_name]
    style.font.name = font_name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    # python-docx's font.name only sets the "western" font slot — Word's docx
    # schema stores separate hint slots that override it during rendering.
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), font_name)


def _add_heading_bottom_border(document: Document, style_name: str, color: str, size: int = 12):
    """Adds a bottom border under every paragraph using this style — used for
    the Corporate Bold template's underline bar beneath section headings."""

    style = document.styles[style_name]
    ppr = style.element.get_or_add_pPr()
    pbdr = ppr.makeelement(qn("w:pBdr"), {})
    bottom = pbdr.makeelement(qn("w:bottom"), {
        qn("w:val"): "single",
        qn("w:sz"): str(size),
        qn("w:space"): "4",
        qn("w:color"): color,
    })
    pbdr.append(bottom)
    ppr.append(pbdr)


def _build(filename: str, *, heading1, heading2, normal):
    document = Document()
    _set_style(document, "Normal", *normal)
    _set_style(document, "Heading 1", *heading1)
    _set_style(document, "Heading 2", *heading2)
    document.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    document.save(_OUTPUT_DIR / filename)


# The Professional (default) template's palette. Only headings and
# subheadings are coloured — body text, tables and everything else stay
# black, so the DOCX and the PDF can look the same. These two values are
# duplicated in html/template_1.html's CSS (--heading / --subheading); change
# both together or the two export formats drift apart.
HEADING_COLOR = "0D2B5E"
SUBHEADING_COLOR = "1A5FB4"
BODY_FONT = "Arial"


def _build_professional() -> None:
    """Reference doc for html/template_1.html (template_id 1).

    Pandoc maps HTML <h1>../<h4> onto the Heading 1..4 *styles* of this file,
    so defining the colours here is what makes the DOCX export match the PDF —
    Pandoc ignores the CSS in the HTML entirely.
    """

    document = Document()
    _set_style(document, "Normal", BODY_FONT, 10.5, "1A1A1A", False)
    _set_style(document, "Title", BODY_FONT, 22, HEADING_COLOR, True)
    _set_style(document, "Heading 1", BODY_FONT, 17, HEADING_COLOR, True)
    _set_style(document, "Heading 2", BODY_FONT, 13, HEADING_COLOR, True)
    _set_style(document, "Heading 3", BODY_FONT, 11, SUBHEADING_COLOR, True)
    _set_style(document, "Heading 4", BODY_FONT, 10.5, SUBHEADING_COLOR, True)
    # Matches the PDF's rule under each section heading.
    _add_heading_bottom_border(document, "Heading 2", HEADING_COLOR, size=8)
    for style_name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        document.styles[style_name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    document.save(_OUTPUT_DIR / "professional.docx")


def main():
    _build_professional()

    _build(
        "classic_serif.docx",
        normal=("Georgia", 11, "1A1A1A", False),
        heading1=("Georgia", 26, "1F3864", True),
        heading2=("Georgia", 15, "1F3864", True),
    )
    _build(
        "modern_blue.docx",
        normal=("Calibri", 11, "222222", False),
        heading1=("Calibri", 26, "0B5394", True),
        heading2=("Calibri", 14, "0B5394", True),
    )
    _build(
        "minimal_mono.docx",
        normal=("Arial", 10.5, "333333", False),
        heading1=("Arial", 22, "333333", False),
        heading2=("Arial", 13, "666666", True),
    )

    document = Document()
    _set_style(document, "Normal", "Arial", 11, "1A1A1A", False)
    _set_style(document, "Heading 1", "Arial", 28, "003366", True)
    _set_style(document, "Heading 2", "Arial", 15, "003366", True)
    _add_heading_bottom_border(document, "Heading 1", "003366", size=18)
    document.save(_OUTPUT_DIR / "corporate_bold.docx")

    print("Generated 5 reference .docx templates in", _OUTPUT_DIR)


if __name__ == "__main__":
    main()
