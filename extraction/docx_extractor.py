from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from extraction.base import BaseExtractor, ExtractedDocument, ExtractedPage, ExtractionMethod
from extraction.heading_detector import LineFeatures, classify_heading

HEADING_STYLE_TO_MARKDOWN = {
    "Heading 1": "#",
    "Heading 2": "##",
    "Heading 3": "###",
    "Title": "#",
}


class DocxExtractor(BaseExtractor):
    """python-docx based conversion to Markdown, preserving heading levels,
    tables, and lists. DOCX has no reliable page boundaries, so the whole
    document is emitted as a single ExtractedPage."""

    def extract(self, file_path: str, source_filename: str) -> ExtractedDocument:
        document = Document(file_path)
        body_size = self._estimate_body_font_size(document)
        markdown_parts: list[str] = []

        for block in self._iter_block_items(document):
            if isinstance(block, Paragraph):
                rendered = self._render_paragraph(block, body_size)
            else:
                rendered = self._render_table(block)
            if rendered:
                markdown_parts.append(rendered)

        markdown = "\n\n".join(markdown_parts).strip()
        print(markdown)
        page = ExtractedPage(page_number=1, markdown=markdown, extraction_method=ExtractionMethod.DOCX)
        return ExtractedDocument(markdown=markdown, source_filename=source_filename, pages=[page])

    # ------------------------------------------------------------------
    # Document-order traversal — python-docx exposes paragraphs and tables
    # as separate collections, so this walks the underlying XML body to
    # preserve their original interleaving.
    # ------------------------------------------------------------------

    def _iter_block_items(self, document: Document):
        body = document.element.body
        for child in body.iterchildren():
            if child.tag.endswith("}p"):
                yield Paragraph(child, document)
            elif child.tag.endswith("}tbl"):
                yield Table(child, document)

    def _render_paragraph(self, paragraph: Paragraph, body_size: float) -> str:
        text = paragraph.text.strip()
        if not text:
            return ""

        style_name = paragraph.style.name if paragraph.style else ""

        prefix = HEADING_STYLE_TO_MARKDOWN.get(style_name)
        if prefix:
            return f"{prefix} {text}"

        if style_name.startswith("List Bullet"):
            return f"- {text}"
        if style_name.startswith("List Number"):
            return f"1. {text}"

        fallback_prefix = self._detect_heading_by_formatting(paragraph, body_size)
        if fallback_prefix:
            return f"{fallback_prefix} {text}"

        return text

    def _estimate_body_font_size(self, document: Document) -> float:
        """Body text is assumed to be the most common font size (by character
        count) across the document — headings are detected relative to this
        baseline, mirroring PDFExtractor's approach."""

        size_counts: dict[float, int] = {}
        for paragraph in document.paragraphs:
            for run in paragraph.runs:
                if run.font.size is None:
                    continue
                text_len = len(run.text.strip())
                if text_len:
                    size_pt = round(run.font.size.pt, 1)
                    size_counts[size_pt] = size_counts.get(size_pt, 0) + text_len

        if not size_counts:
            return 11.0
        return max(size_counts, key=lambda size: size_counts[size])

    def _detect_heading_by_formatting(self, paragraph: Paragraph, body_size: float):
        """Fallback for documents that mark section titles via direct bold/
        underline/font-size formatting on a "Normal"-styled paragraph instead
        of named Word heading styles (common in copy-pasted or informally
        authored RFPs). Bold/underline are measured as a share of the
        paragraph's characters — not "any run is bold" — so a bolded bullet
        lead-in ("**Content Ingestion:** the tool must...") scores low and
        isn't mistaken for a heading, while a fully-bold short line scores high."""

        runs_with_text = [run for run in paragraph.runs if run.text.strip()]
        if not runs_with_text:
            return None

        total_chars = sum(len(run.text) for run in runs_with_text)
        bold_chars = sum(len(run.text) for run in runs_with_text if run.bold)
        underline_chars = sum(len(run.text) for run in runs_with_text if run.underline)

        sizes = [run.font.size.pt for run in runs_with_text if run.font.size is not None]
        max_size = max(sizes) if sizes else body_size

        return classify_heading(LineFeatures(
            text=paragraph.text,
            size=max_size,
            body_size=body_size,
            bold_ratio=bold_chars / total_chars if total_chars else 0.0,
            underline_ratio=underline_chars / total_chars if total_chars else 0.0,
        ))

    def _render_table(self, table: Table) -> str:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        rows = [row for row in rows if any(cell for cell in row)]
        if not rows:
            return ""

        header, *body_rows = rows
        lines = [
            "| " + " | ".join(header) + " |",
            "| " + " | ".join("---" for _ in header) + " |",
        ]
        lines.extend("| " + " | ".join(row) + " |" for row in body_rows)
        return "\n".join(lines)

if __name__ == "__main__":
    obj1 = DocxExtractor()
    obj1.extract("/home/ib-40/Downloads/RFP_Inmar.docx", "knowledge document.docx")